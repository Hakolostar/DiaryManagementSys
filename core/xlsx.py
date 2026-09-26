"""Excel (.xlsx) export helper built on openpyxl."""

from io import BytesIO

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="047857")  # emerald-700
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)


def build_xlsx_response(filename, headers, rows):
    """Return an HttpResponse that downloads `headers`/`rows` as a .xlsx file."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Export"

    for col, head in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=head)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for r, row in enumerate(rows, start=2):
        for c, val in enumerate(row, start=1):
            if isinstance(val, bool):
                val = "Yes" if val else "No"
            ws.cell(row=r, column=c, value="" if val is None else val)

    ws.freeze_panes = "A2"
    for col in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 22

    buf = BytesIO()
    wb.save(buf)
    response = HttpResponse(
        buf.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}.xlsx"'
    return response